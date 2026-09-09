# -*- coding: utf-8 -*-
"""Final Word: 校辩智生 演示环境部署与使用说明."""
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(r"e:\zzs\github\AgentA3\docs\deployment")
FIG = ROOT / "figures"
OUT = ROOT / "校辩智生-演示环境部署与使用说明.docx"
DESKTOP_DIR = Path(r"c:\Users\zzs\Desktop\视频剪辑\文档")
DESKTOP = DESKTOP_DIR / "校辩智生-演示环境部署与使用说明.docx"
DESKTOP2 = Path(r"c:\Users\zzs\Desktop\校辩智生-演示环境部署与使用说明.docx")


def set_run_font(run, name="宋体", size=12, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    ascii_font = "Times New Roman" if name != "黑体" else "黑体"
    r_fonts.set(qn("w:ascii"), ascii_font)
    r_fonts.set(qn("w:hAnsi"), ascii_font)
    r_fonts.set(qn("w:eastAsia"), name)


def add_h(doc, text, level=1):
    p = doc.add_paragraph()
    size = {0: 18, 1: 14, 2: 12}.get(level, 12)
    set_run_font(p.add_run(text), "黑体", size, True)
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5


def add_p(doc, text, indent=True):
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), "宋体", 12)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)


def add_bullet(doc, text):
    p = doc.add_paragraph()
    set_run_font(p.add_run("• " + text), "宋体", 12)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.line_spacing = 1.5


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(text), "宋体", 10.5)
    p.paragraph_format.space_after = Pt(10)


def add_fig(doc, name, caption, width_cm=14.5):
    path = FIG / name
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            set_run_font(p.add_run(val), "宋体", 10.5, bold=(i == 0))
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "E7E6E6")
                shd.set(qn("w:val"), "clear")
                cell._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph()


def add_code(doc, code):
    for line in code.splitlines() or [""]:
        p = doc.add_paragraph()
        set_run_font(p.add_run(line if line else " "), "Consolas", 9)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F5F5F5")
        shd.set(qn("w:val"), "clear")
        p._element.get_or_add_pPr().append(shd)
    doc.add_paragraph()


def main():
    doc = Document()

    add_h(doc, "校辩智生——演示环境部署与使用说明", 0)
    add_p(doc, "终版｜2026-09-05｜用于评委访问、账号登录与 APP 演示。", indent=False)

    # ===== 最开始：访问与账号 =====
    add_h(doc, "一、访问地址", 1)
    add_table(
        doc,
        [
            ["入口", "地址"],
            ["Web 管理端登录", "http://129.211.82.112:3000/"],
            ["智能体设置", "http://129.211.82.112:3000/ai/agent-settings"],
            ["API（同域代理）", "http://129.211.82.112:3000/api/**"],
        ],
    )
    add_p(doc, "浏览器打开上述地址即可访问。Web 端通过 Nginx 将 /api 代理到后端，无需单独记忆后端端口。")

    add_h(doc, "二、测试账号", 1)
    add_h(doc, "2.1 APP / 小程序（学生端）", 2)
    add_table(
        doc,
        [
            ["登录端", "用户名", "密码", "角色"],
            ["APP / 小程序", "zzs", "Zhang@2114", "STUDENT"],
        ],
    )
    add_p(doc, "APP 安装包随附件发送，安装后使用上表账号登录。")

    add_h(doc, "2.2 Web 管理端", 2)
    add_table(
        doc,
        [
            ["登录端", "用户名", "密码", "角色", "用途"],
            ["Web 管理端", "admin", "admin123", "ADMIN", "智能体、知识库、题库等治理"],
            ["Web 管理端", "merchant01", "admin123", "MERCHANT", "商家相关页面"],
        ],
    )

    add_h(doc, "2.3 登录说明", 2)
    add_bullet(doc, "Web 管理端仅支持 ADMIN、MERCHANT；请使用 admin / admin123。")
    add_bullet(doc, "APP / 小程序仅支持 STUDENT、TEACHER；演示请使用 zzs / Zhang@2114。")
    add_bullet(doc, "不要用 APP 账号登录 Web，也不要用 admin 登录 APP。")

    add_fig(doc, "deploy-01-web-login.png", "图1  Web管理端登录页", 14.0)
    add_fig(doc, "deploy-02-agent-settings-live.png", "图2  智能体设置页", 14.0)

    # ===== APP =====
    add_h(doc, "三、APP 安装与使用", 1)
    add_p(doc, "Android 安装包已随附件发送，请直接安装附件中的 APK。安装完成后打开应用，使用账号 zzs、密码 Zhang@2114 登录。")
    add_p(doc, "登录后可体验校园服务、AI 助手、个性化学习、题库考试、会议等功能。")
    add_fig(doc, "deploy-04-app-home.png", "图3  APP首页", 7.0)
    add_fig(doc, "deploy-05-app-ai.png", "图4  AI智能助手", 7.0)

    # ===== Web 验收 =====
    add_h(doc, "四、Web 管理端使用", 1)
    add_p(doc, "1. 打开 http://129.211.82.112:3000/ ，使用 admin / admin123 登录。")
    add_p(doc, "2. 进入「AI 模块 → 智能体设置」：http://129.211.82.112:3000/ai/agent-settings")
    add_p(doc, "3. 可继续查看：模型配置、知识库管理、题库/试卷、用户与角色、会议语音模型配置等。")
    add_fig(doc, "deploy-03-users.png", "图5  用户与角色管理", 14.0)

    # ===== 部署复现（精简，无占位）=====
    add_h(doc, "五、服务器部署摘要（可选复现）", 1)
    add_p(doc, "系统由 Web 管理端、Java 后端、Python AI 服务、MySQL、Redis 组成，提交清单为 deploy/compose.submission.yml。")
    add_code(
        doc,
        """cp deploy/.env.example deploy/.env
# 填写 MYSQL_ROOT_PASSWORD、JWT_SECRET、AI_INTERNAL_TOKEN
docker compose --env-file deploy/.env -f deploy/compose.submission.yml up -d --build
bash deploy/verify.sh""",
    )
    add_p(doc, "更完整的 Compose 与排障说明见 docs/deployment/submission-runbook.md。")

    add_h(doc, "六、常见问题", 1)
    add_table(
        doc,
        [
            ["现象", "处理"],
            ["Web 登录失败", "确认使用 admin / admin123，且从 Web 入口登录"],
            ["APP 登录失败", "确认使用 zzs / Zhang@2114，且网络可访问部署地址"],
            ["AI 无回复", "检查智能体设置页工具开关与模型绑定，确认 AI 服务正常"],
            ["页面能开但接口 502", "检查后端容器状态与 /api 反代"],
        ],
    )

    add_h(doc, "七、说明", 1)
    add_p(doc, "本文为终版演示说明。账号仅用于比赛/答辩演示；正式生产环境须另行加固口令与密钥管理。")

    doc.save(str(OUT))
    print("saved", OUT, OUT.stat().st_size)

    md = ROOT / "校辩智生-演示环境部署与使用说明.md"
    md.write_text(
        """# 校辩智生——演示环境部署与使用说明（终版）

## 一、访问地址

| 入口 | 地址 |
|---|---|
| Web 管理端登录 | http://129.211.82.112:3000/ |
| 智能体设置 | http://129.211.82.112:3000/ai/agent-settings |

## 二、测试账号

### APP / 小程序
| 用户名 | 密码 | 角色 |
|---|---|---|
| zzs | Zhang@2114 | STUDENT |

APP 安装包随附件发送。

### Web 管理端
| 用户名 | 密码 | 角色 |
|---|---|---|
| admin | admin123 | ADMIN |
| merchant01 | admin123 | MERCHANT |

Word 终版：`校辩智生-演示环境部署与使用说明.docx`
""",
        encoding="utf-8",
    )

    for dest in (DESKTOP, DESKTOP2, ROOT / "演示环境部署与使用说明.docx"):
        try:
            DESKTOP_DIR.mkdir(parents=True, exist_ok=True)
            shutil.copy2(OUT, dest)
            print("copied", dest)
        except Exception as exc:
            print("copy fail", dest, exc)


if __name__ == "__main__":
    main()
