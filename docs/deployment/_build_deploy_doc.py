# -*- coding: utf-8 -*-
"""Build deployment guide DOCX with screenshots."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"e:\zzs\github\AgentA3\docs\deployment")
FIG = ROOT / "figures"
OUT = ROOT / "演示环境部署与使用说明.docx"
DESKTOP = Path(r"c:\Users\zzs\Desktop\视频剪辑\文档\演示环境部署与使用说明.docx")


def set_run_font(run, name="宋体", size=12, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    ascii_font = "Times New Roman" if name != "黑体" else "黑体"
    r_fonts.set(qn("w:ascii"), ascii_font)
    r_fonts.set(qn("w:hAnsi"), ascii_font)
    r_fonts.set(qn("w:eastAsia"), name)


def add_h(doc, text, level=1):
    p = doc.add_paragraph()
    size = {0: 18, 1: 14, 2: 12, 3: 12}[level]
    set_run_font(p.add_run(text), "黑体", size, True)
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5


def add_p(doc, text, indent=True):
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), "宋体", 12)
    p.paragraph_format.line_spacing = 1.5
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)


def add_bullet(doc, text):
    p = doc.add_paragraph()
    set_run_font(p.add_run("• " + text), "宋体", 12)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.line_spacing = 1.5


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(text), "宋体", 10.5)
    p.paragraph_format.space_after = Pt(10)


def add_fig(doc, name, caption, width_cm=14.5):
    path = FIG / name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if path.exists():
        p.add_run().add_picture(str(path), width=Cm(width_cm))
    else:
        set_run_font(p.add_run(f"[待补图: {name}]"), "宋体", 10.5)
        # gray box
        pass
    add_caption(doc, caption)


def add_placeholder(doc, caption, hint):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(caption), "宋体", 10.5, True)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Cm(14)
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9D9D9")
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)
    cell.text = ""
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cp.add_run("【截图占位】\n" + hint), "宋体", 10.5)
    for _ in range(3):
        cell.add_paragraph()
    doc.add_paragraph()


def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            set_run_font(p.add_run(val), "宋体", 9.5 if cols > 4 else 10.5, bold=(i == 0))
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "E7E6E6")
                shd.set(qn("w:val"), "clear")
                cell._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph()


def add_code(doc, code):
    for line in code.splitlines() or [""]:
        p = doc.add_paragraph()
        set_run_font(p.add_run(line if line else " "), "Consolas", 9)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F5F5F5")
        shd.set(qn("w:val"), "clear")
        p._element.get_or_add_pPr().append(shd)
    doc.add_paragraph()


def main():
    doc = Document()

    add_h(doc, "AgentA3 演示环境部署与使用说明", 0)
    add_p(doc, "版本 V1.1｜更新日期 2026-09-05｜面向评委复现、答辩演示与团队联调。", indent=False)

    add_h(doc, "1. 已部署线上地址", 1)
    add_table(
        doc,
        [
            ["入口", "地址", "说明"],
            ["Web 登录 / 首页", "http://129.211.82.112:3000/", "智慧校园控制台"],
            ["智能体设置", "http://129.211.82.112:3000/ai/agent-settings", "AI 治理核心页"],
            ["API 代理", "http://129.211.82.112:3000/api/**", "Nginx 反代到 Java"],
        ],
    )
    add_p(doc, "线上 Web 使用相对路径模式，浏览器只访问 3000 端口即可。")
    add_fig(doc, "deploy-01-web-login.png", "图1  Web管理端登录页（线上）", 14.5)
    add_fig(doc, "deploy-02-agent-settings-live.png", "图2  智能体设置页（线上）", 14.5)

    add_h(doc, "2. 测试账号（请按此表登录）", 1)
    add_p(doc, "种子数据来自 AppBackend/src/main/resources/data.sql 与 AppBackend/02_test_users.sql。当前演示环境统一密码为 admin123（仅比赛演示，非生产安全基线）。")

    add_h(doc, "2.1 推荐答辩账号", 2)
    add_table(
        doc,
        [
            ["用途", "登录端", "用户名", "密码", "角色"],
            ["管理治理", "Web 管理端", "admin", "admin123", "ADMIN"],
            ["学生学习主线", "小程序 / APP", "zzs", "admin123", "STUDENT"],
            ["教师能力", "小程序 / APP", "fjj", "admin123", "TEACHER"],
            ["商家能力", "Web 管理端", "merchant01", "admin123", "MERCHANT"],
        ],
    )

    add_h(doc, "2.2 补充测试账号", 2)
    add_table(
        doc,
        [
            ["用户名", "密码", "角色", "登录端", "备注"],
            ["test_admin", "admin123", "ADMIN", "Web", "标准测试管理员"],
            ["test_teacher", "admin123", "TEACHER", "小程序/APP", "标准测试教师"],
            ["test_student", "admin123", "STUDENT", "小程序/APP", "标准测试学生"],
            ["test_merchant", "admin123", "MERCHANT", "Web", "标准测试商家"],
            ["qb_peer", "admin123", "STUDENT", "小程序/APP", "题库公私对照"],
            ["lisi / wangwu 等", "admin123", "STUDENT", "小程序/APP", "普通学生样例"],
            ["merchant02~04", "admin123", "MERCHANT", "Web", "其他商家样例"],
        ],
    )

    add_h(doc, "2.3 登录边界", 2)
    add_table(
        doc,
        [
            ["入口", "允许角色", "拒绝角色"],
            ["Web weblogin", "ADMIN、MERCHANT", "STUDENT、TEACHER"],
            ["App/小程序 applogin", "STUDENT、TEACHER", "ADMIN、MERCHANT"],
        ],
    )
    add_bullet(doc, "不要用 zzs 登录 Web；不要用 admin 登录小程序/APP。")
    add_bullet(doc, "智能体设置页右上角「测试用户」可填 zzs，用于管理端代测学生侧 AI 能力。")

    add_h(doc, "3. APP / 小程序端", 1)
    add_p(doc, "移动端工程为 mini_program_app（uni-app）。本次演示不使用 Frontend/ 目录。")
    add_h(doc, "3.1 指向线上后端", 2)
    add_code(
        doc,
        """# mini_program_app/.env.production.local（不提交 Git）
VITE_API_BASE_URL=http://129.211.82.112:3000""",
    )
    add_h(doc, "3.2 HBuilderX 导出", 2)
    add_bullet(doc, "打开 mini_program_app → 发行 → 原生 App-云打包（生成 APK）。")
    add_bullet(doc, "或发行 → 小程序-微信（需 AppID；体验版上传微信后台）。")
    add_placeholder(doc, "图3  HBuilderX 发行/云打包界面", "请截取发行成功页，保存为 deploy-06-hbuilder-export.png")
    add_placeholder(doc, "图4  APP 安装包二维码或下载页", "请放置 APK 下载二维码，保存为 deploy-07-app-qr.png")

    add_h(doc, "3.3 APP 安装包交接表", 2)
    add_table(
        doc,
        [
            ["项", "填写"],
            ["包名 / 文件名", ""],
            ["存放位置 / 下载链接", ""],
            ["构建时间 / Git SHA", ""],
            ["最低 Android 版本", ""],
        ],
    )
    add_p(doc, "安装后使用 zzs / admin123 登录。")
    add_fig(doc, "deploy-04-app-home.png", "图5  APP/小程序首页", 7.2)
    add_fig(doc, "deploy-05-app-ai.png", "图6  AI智能助手界面", 7.2)

    add_h(doc, "4. Web 管理端快速验收", 1)
    add_bullet(doc, "打开 http://129.211.82.112:3000/ ，使用 admin / admin123 登录。")
    add_bullet(doc, "进入 AI 模块 → 智能体设置，检查工具开关、模型绑定、题库映射。")
    add_bullet(doc, "再抽查：模型配置、知识库、题库/试卷、用户与角色、会议语音模型配置。")
    add_fig(doc, "deploy-03-users.png", "图7  用户与角色管理", 14.5)

    add_h(doc, "5. 服务器一键部署（从零复现）", 1)
    add_p(doc, "细节见 docs/deployment/submission-runbook.md 与 AppBackend/deploy/README.md。")
    add_code(
        doc,
        """cp deploy/.env.example deploy/.env
# 填写 MYSQL_ROOT_PASSWORD、JWT_SECRET、AI_INTERNAL_TOKEN
# PUBLIC_BASE_URL=http://129.211.82.112:3000

docker compose --env-file deploy/.env -f deploy/compose.submission.yml up -d --build
bash deploy/verify.sh

# 或服务器脚本
./AppBackend/deploy/deploy-on-server.sh""",
    )

    add_h(doc, "6. 常见问题", 1)
    add_table(
        doc,
        [
            ["现象", "处理"],
            ["登录失败", "核对角色与入口；密码是否 admin123；库中是否有用户"],
            ["小程序请求失败", "检查 VITE_API_BASE_URL 是否指向线上 3000"],
            ["AI 无回复", "检查智能体设置与 AI 容器健康、模型绑定"],
            ["/api 502", "查看 Java 容器日志与 Compose 网络"],
        ],
    )

    add_h(doc, "7. 安全提示", 1)
    add_p(doc, "本文账号仅用于比赛演示。正式环境须改用密码哈希、轮换密钥，并禁止在截图中暴露真实云密钥与内部令牌。")

    doc.save(str(OUT))
    print("saved", OUT, OUT.stat().st_size)
    try:
        import shutil

        shutil.copy2(OUT, DESKTOP)
        print("copied", DESKTOP)
    except Exception as exc:
        print("desktop copy failed", exc)
        alt = Path(r"c:\Users\zzs\Desktop\演示环境部署与使用说明.docx")
        try:
            shutil.copy2(OUT, alt)
            print("copied alt", alt)
        except Exception as e2:
            print("alt failed", e2)


if __name__ == "__main__":
    main()
