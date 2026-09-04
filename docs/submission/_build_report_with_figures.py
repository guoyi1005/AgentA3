# -*- coding: utf-8 -*-
"""Build complete AgentA3 work report DOCX with all screenshots."""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(r"e:\zzs\github\AgentA3\docs\submission")
FIG = ROOT / "figures"
OUT = ROOT / "作品报告-AgentA3.docx"
OUT_DESKTOP = Path(r"c:\Users\zzs\Desktop\视频剪辑\文档\作品报告-AgentA3.docx")
TPL = Path(r"c:\Users\zzs\Desktop\视频剪辑\文档\05-3 作品报告（人工智能挑战赛，2023版）模板.docx")


def set_run_font(run, name="宋体", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    ascii_font = "Consolas" if name == "Consolas" else ("Times New Roman" if name != "黑体" else "黑体")
    r_fonts.set(qn("w:ascii"), ascii_font)
    r_fonts.set(qn("w:hAnsi"), ascii_font)
    r_fonts.set(qn("w:eastAsia"), "宋体" if name == "Consolas" else name)
    if color:
        run.font.color.rgb = color


def p_style(p, before=0, after=6, line=1.5, indent=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Cm(0.74)


def add_title(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(text), "黑体", 18, True)
    elif level == 1:
        set_run_font(p.add_run(text), "黑体", 14, True)
        p_style(p, before=14, after=8)
    elif level == 2:
        set_run_font(p.add_run(text), "黑体", 12, True)
        p_style(p, before=10, after=6)
    else:
        set_run_font(p.add_run(text), "黑体", 12, True)
        p_style(p, before=8, after=4)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), "宋体", 12)
    p_style(p, indent=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    set_run_font(p.add_run("• " + text), "宋体", 12)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(text), "宋体", 10.5)
    p_style(p, before=2, after=10, line=1.2)
    return p


def add_figure(doc, filename, caption, width_cm=7.2):
    path = FIG / filename
    if not path.exists():
        add_body(doc, f"[缺图: {filename}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    p_style(p, before=6, after=2, line=1.0)
    add_caption(doc, caption)


def add_two_figures(doc, items, width_cm=7.0):
    """items: list of (filename, caption), length 1 or 2."""
    table = doc.add_table(rows=2, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (filename, caption) in enumerate(items):
        path = FIG / filename
        cell0 = table.cell(0, j)
        cell0.text = ""
        p = cell0.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if path.exists():
            run = p.add_run()
            run.add_picture(str(path), width=Cm(width_cm))
        cell1 = table.cell(1, j)
        cell1.text = ""
        cp = cell1.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(cp.add_run(caption), "宋体", 9)
    doc.add_paragraph()


def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            set_run_font(p.add_run(val), "宋体", 10, bold=(i == 0))
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "E7E6E6")
                shd.set(qn("w:val"), "clear")
                cell._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph()


def add_code(doc, code):
    for line in code.splitlines() or [""]:
        p = doc.add_paragraph()
        set_run_font(p.add_run(line if line else " "), "Consolas", 8.5)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F5F5F5")
        shd.set(qn("w:val"), "clear")
        p._element.get_or_add_pPr().append(shd)
    doc.add_paragraph()


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build():
    try:
        doc = Document(str(TPL))
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
    except Exception:
        doc = Document()

    # ===== Cover =====
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("中国计算机学会青少年\n人工智能挑战赛\n作品报告"), "黑体", 20, True)

    for _ in range(2):
        doc.add_paragraph()
    for k, v in [
        ("作品编号：", "（填写）"),
        ("作品名称：", "AgentA3——智慧校园个性化学习多智能体系统"),
        ("编写日期：", "2026-09-05"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(k), "宋体", 14, True)
        set_run_font(p.add_run(v), "宋体", 14)

    page_break(doc)

    # ===== TOC-like outline =====
    add_title(doc, "目  录", 0)
    for item in [
        "第1章  作品概述",
        "第2章  任务分析",
        "第3章  原理介绍",
        "第4章  系统实现",
        "第5章  测试分析",
        "第6章  作品总结",
    ]:
        p = doc.add_paragraph()
        set_run_font(p.add_run(item), "宋体", 12)
        p_style(p, after=8)

    page_break(doc)

    # ==================== Ch1 ====================
    add_title(doc, "第1章  作品概述", 1)
    add_title(doc, "1.1 建设背景与目标", 2)
    add_body(
        doc,
        "高等教育场景中，学生的课程学习、研讨会议、练习考试与校园日常服务往往分散在多个入口。"
        "通用大模型虽能生成内容，但缺少可持续沉淀的学习背景，也难以把过程信号转化为可治理的教学证据。"
        "AgentA3 面向上述问题，构建以“证据驱动的个性化学习闭环”为主线的多端系统：将对话补问、资源互动、"
        "会议转写与客观考试结果转化为具备来源、置信度与状态的证据；由 Java 维护七维画像与学习路径；"
        "由 Python Leader 路由专业智能体，或进入 Python 课程受控资源生成工作流，最终交付可复审、可导出的多模态学习资源。",
    )
    add_body(doc, "作品目标可概括为三点：可追踪（画像/资源/考试均可回溯）、可协作（校园 Leader 与课程 DAG 双拓扑隔离）、可落地（保留校园服务并增量接入学习智能体能力）。")

    add_title(doc, "1.2 作品定位", 2)
    add_body(
        doc,
        "系统将学生的对话、研讨和资源互动转化为可追溯证据，再由 Leader 编排专业智能体，"
        "为不同学生生成并推送不同的学习资源。创新重点不在“再做一个聊天机器人”，而在工程边界与可信机制。",
    )

    add_title(doc, "1.3 系统组成", 2)
    add_table(
        doc,
        [
            ["工程边界", "技术基线", "主要职责"],
            ["移动端小程序", "uni-app（mini_program_app）", "学生/教师：校园服务、AI 助手、学习、会议、考试"],
            ["Web 管理端", "React/Vite（AppWeb）", "管理员：知识库、智能体、题库组卷、画像规则、校园业务治理"],
            ["Java 业务后端", "Spring Boot（AppBackend）", "身份权限、画像规则、工作流编排、题库考试、会议 ASR 桥接"],
            ["Python AI 服务", "FastAPI（ai-servers）", "Leader/Catalog、课程 typed DAG、模型调用、资源复审导出"],
        ],
    )

    add_two_figures(
        doc,
        [
            ("app-01-home.png", "图1-1  小程序首页门户"),
            ("app-21-ai-create.png", "图1-2  AI创作与校园课程入口"),
        ],
    )
    add_figure(doc, "admin-01-users.png", "图1-3  Web管理端（用户与角色治理）", width_cm=14.5)

    add_title(doc, "1.4 预期效果", 2)
    add_body(doc, "对学生：在保留校园生活入口的同时，获得画像补问→资源生成→路径/推荐→练习考试→掌握度反馈的完整链路。")
    add_body(doc, "对管理员：在 Web 端配置知识库、模型与智能体，生成并治理题库试卷，审计会议与画像规则。")
    add_body(doc, "对评审：关键结论可回到生产代码与界面证据，不以未实测指标替代工程事实。")

    page_break(doc)

    # ==================== Ch2 ====================
    add_title(doc, "第2章  任务分析", 1)
    add_title(doc, "2.1 任务拆解", 2)
    add_table(
        doc,
        [
            ["任务", "赛题关注点", "本作品对应实现", "状态"],
            ["T1", "对话式学习画像（≥6维）", "Java七维画像+证据协议；小程序雷达页", "已实现"],
            ["T2", "多智能体资源生成（≥5类）", "Catalog+课程DAG（讲解/导图/练习/代码/拓展）+PPT管线", "已实现"],
            ["T3", "个性化路径与推送", "学习计划拆解、路径规划与推荐", "已实现"],
            ["T4", "智能辅导（加分）", "Leader助手、图表/写作/编程辅导", "已实现"],
            ["T5", "学习效果评估（加分）", "题库组卷、在线编程题、客观题反馈闭环", "已实现"],
        ],
    )

    add_title(doc, "2.2 目标用户与角色", 2)
    add_table(
        doc,
        [
            ["角色", "入口", "核心诉求"],
            ["STUDENT", "小程序 applogin", "个性化学习、会议、考试、校园服务"],
            ["TEACHER", "小程序 applogin", "学习与会议相关已授权能力"],
            ["ADMIN", "Web weblogin", "MaxKB、题库、智能体、画像规则与校园业务治理"],
            ["MERCHANT", "Web weblogin", "商户业务；教学治理接口仍受ADMIN约束"],
        ],
    )

    add_title(doc, "2.3 小程序端主要功能", 2)
    add_body(doc, "小程序在保留校园服务的基础上，形成 AI 学习与创作能力矩阵。主要功能模块如下。")

    add_title(doc, "2.3.1 校园服务", 3)
    add_body(doc, "覆盖活动、闲置交易、论坛、优惠券、校园地图、课表等日常入口，保证系统可落地而非仅演示 AI。")
    add_two_figures(doc, [("app-02-activities.png", "图2-1  校园活动"), ("app-03-secondhand.png", "图2-2  闲置交易")])
    add_two_figures(doc, [("app-04-forum.png", "图2-3  校园论坛"), ("app-05-coupons.png", "图2-4  优惠券")])
    add_two_figures(doc, [("app-06-map.png", "图2-5  校园地图"), ("app-23-schedule.png", "图2-6  我的课表")])

    add_title(doc, "2.3.2 AI 助手与多模态创作", 3)
    add_body(doc, "智能助手支持会话历史、工具调用与可视化产物（流程图/思维导图等）；同时提供智能写作、PPT、架构图、流程图、思维导图等专属入口。")
    add_two_figures(doc, [("app-27-ai-assistant.png", "图2-7  智能助手对话"), ("app-25-ai-history.png", "图2-8  AI会话历史")])
    add_two_figures(doc, [("app-10-smart-writing.png", "图2-9  智能写作"), ("app-16-mindmap.png", "图2-10  AI思维导图")])
    add_two_figures(doc, [("app-17-architecture.png", "图2-11  AI架构图"), ("app-18-flowchart.png", "图2-12  AI流程图")])
    add_two_figures(doc, [("app-11-ppt-template.png", "图2-13  PPT模板选择"), ("app-12-ppt-done.png", "图2-14  PPT生成完成")])

    add_title(doc, "2.3.3 个性化学习、题库与考试", 3)
    add_body(doc, "包括个人画像雷达、学习计划拆解、我的课程、题库生成、试卷选题、编程题库等，支撑“学—练—测”闭环。")
    add_two_figures(doc, [("app-26-profile-radar.png", "图2-15  个人画像雷达"), ("app-13-plan-decompose.png", "图2-16  学习计划拆解")])
    add_two_figures(doc, [("app-24-my-courses.png", "图2-17  我的课程"), ("app-14-question-gen.png", "图2-18  题库生成")])
    add_two_figures(doc, [("app-15-paper-select.png", "图2-19  试卷选题"), ("app-19-coding-bank.png", "图2-20  编程题库")])

    add_title(doc, "2.3.4 智能会议", 3)
    add_body(doc, "支持会议首页、AI 会议纪要与任务/转写记录，服务端桥接实时 ASR，会后由会议类智能体生成纪要。")
    add_two_figures(doc, [("app-07-meeting-home.png", "图2-21  会议首页"), ("app-08-meeting-summary.png", "图2-22  AI会议纪要")])
    add_figure(doc, "app-09-meeting-detail.png", "图2-23  会议详情与任务记录", width_cm=7.2)

    add_title(doc, "2.3.5 学习辅助工具", 3)
    add_body(doc, "另提供 PDF→Word、水印工具等实用能力，降低学习资料处理成本。")
    add_two_figures(doc, [("app-20-pdf2word.png", "图2-24  PDF转Word"), ("app-22-watermark.png", "图2-25  水印工具")])

    add_title(doc, "2.4 Web 管理端主要功能", 2)
    add_body(doc, "管理端承担校园业务治理与 AI 运维两类职责：前者覆盖用户、活动、论坛、食堂、地图、闲置、商家等；后者覆盖模型配置、智能体开关/绑定/测试、知识库、画像规则、题库与观测。")

    add_title(doc, "2.4.1 校园业务治理", 3)
    add_figure(doc, "admin-02-activities.png", "图2-26  活动管理", width_cm=14.5)
    add_figure(doc, "admin-03-forum.png", "图2-27  论坛帖子管理", width_cm=14.5)
    add_two_figures(
        doc,
        [("admin-04-canteen.jpg", "图2-28  食堂管理"), ("admin-05-map.jpg", "图2-29  校园标点管理")],
        width_cm=7.0,
    )
    add_figure(doc, "admin-06-secondhand.png", "图2-30  闲置物品管理", width_cm=14.5)
    add_figure(doc, "admin-07-merchants.png", "图2-31  商家管理", width_cm=14.5)

    add_title(doc, "2.4.2 AI 与教学治理", 3)
    add_figure(doc, "admin-21-knowledge.png", "图2-32  AI知识库管理", width_cm=14.5)
    add_figure(doc, "admin-22-profile-rules.png", "图2-33  画像维度规则配置", width_cm=14.5)
    add_figure(doc, "admin-13-model-config.png", "图2-34  AI模型配置", width_cm=14.5)
    add_figure(doc, "admin-14-agent-tools.png", "图2-35  智能体工具开关", width_cm=14.5)
    add_figure(doc, "admin-16-agent-bind.png", "图2-36  智能体模型绑定", width_cm=14.5)
    add_figure(doc, "admin-18-agent-test.png", "图2-37  智能体测试调试", width_cm=14.5)
    add_figure(doc, "admin-11-question-bank.png", "图2-38  题库题目管理", width_cm=14.5)
    add_figure(doc, "admin-10-python-problems.png", "图2-39  Python题库管理", width_cm=14.5)
    add_figure(doc, "admin-09-courses.png", "图2-40  校园课程管理", width_cm=14.5)
    add_figure(doc, "admin-08-meeting-asr.png", "图2-41  会议语音模型配置", width_cm=14.5)
    add_figure(doc, "admin-19-tool-monitor.png", "图2-42  工具调用监控", width_cm=14.5)
    add_figure(doc, "admin-20-observability.png", "图2-43  Langfuse观测配置", width_cm=14.5)
    add_figure(doc, "admin-12-ai-chat.png", "图2-44  管理端智能助手", width_cm=14.5)
    add_figure(doc, "admin-15-agent-formats.png", "图2-45  智能体文件格式设置", width_cm=14.5)
    add_figure(doc, "admin-17-question-mapping.png", "图2-46  智能体题型映射", width_cm=14.5)

    add_title(doc, "2.5 功能交互逻辑", 2)
    add_body(
        doc,
        "小程序与管理端通过 REST/SSE/WebSocket 访问 Spring Boot；Java 负责鉴权与业务事实持久化，"
        "再以内部令牌调用 FastAPI 执行 Leader/DAG/模型任务。模型输出是候选结果，画像分数、路径版本、"
        "试卷与考试状态均由 Java 规则落库，保证“生成能力”与“业务可信”分离。",
    )

    add_title(doc, "2.6 非功能与边界", 2)
    add_table(
        doc,
        [
            ["维度", "要求", "本作品做法"],
            ["交互", "流式输出、卡片化资源", "SSE进度；资源信封含类型/来源/完整性"],
            ["防幻觉", "资料不足不编造", "MaxKB检索不足明确提示；资源复审可拒绝重写"],
            ["安全", "凭据不进客户端", "JWT、内部令牌、第三方Key仅服务端"],
            ["已知限制", "诚实披露", "真实知识库导出、在线压测、ASR生产稳定性仍需专项验收"],
        ],
    )

    page_break(doc)

    # ==================== Ch3 ====================
    add_title(doc, "第3章  原理介绍", 1)
    add_title(doc, "3.1 总体技术路线", 2)
    add_body(
        doc,
        "系统采用“业务治理与模型执行分离”的双服务架构：Java 管身份、领域状态、画像规则与外部编排；"
        "Python 管意图路由、专业智能体与课程 typed DAG。校园通用任务走 Leader 单目标路由；"
        "Python 课程资源包走“路径规划→并行生成→统一复审→打包交付”工作流。",
    )

    add_title(doc, "3.2 七维画像证据协议", 2)
    add_body(
        doc,
        "单次对话或模型判断噪声较大。系统将信号规范化为证据并先入候选池；汇总时按来源可靠性、表达清晰度、"
        "重复度、新鲜度、历史一致性加权置信度；正负证据接近则继续等待；融合步长受单次上限与边界 clamp 约束。"
        "七维包括校园行为、专业课程、学习目标、资源偏好、薄弱知识、学习进度、能力表现。画像总结智能体只读解释，无权直接写分。",
    )
    add_two_figures(
        doc,
        [("app-26-profile-radar.png", "图3-1  小程序七维画像雷达"), ("admin-22-profile-rules.png", "图3-2  管理端画像规则")],
        width_cm=7.0,
    )

    add_title(doc, "3.3 Leader 路由与 Catalog", 2)
    add_body(
        doc,
        "Catalog 维护专业智能体元数据与实现包；Leader 按优先级选择：指定智能体→显式导出→能力查询→视觉生成规则"
        "→高置信校园服务快路径→LLM 规划兜底。Python 课程特定 intent 切入 run_learning_workflow。",
    )
    add_figure(doc, "admin-14-agent-tools.png", "图3-3  Catalog智能体工具开关（管理端）", width_cm=14.5)

    add_title(doc, "3.4 课程资源 typed DAG 与 PPT 管线", 2)
    add_body(
        doc,
        "DAG 内映射讲解文档、思维导图、练习题、代码实验、拓展阅读五类资源智能体，最大并行 3；统一复审后打包，"
        "终态可为 completed 或 partial。产品侧第六类课件由独立 PPT 管线（模板选择→生成→排版修复→导出）补齐。",
    )
    add_two_figures(doc, [("app-11-ppt-template.png", "图3-4  PPT模板与流程"), ("app-12-ppt-done.png", "图3-5  PPT生成与导出")])

    add_title(doc, "3.5 MaxKB 检索与引用式回答", 2)
    add_body(
        doc,
        "MaxKB 仅执行 hit-test；Java 抽取 references 并组装 grounded context；再由本系统 Agent 生成 answer 并返回 citations。"
        "资料不足时明确说明，不编造。",
    )
    add_figure(doc, "admin-21-knowledge.png", "图3-6  知识库管理界面", width_cm=14.5)

    add_title(doc, "3.6 会议实时 ASR 与会后智能体链", 2)
    add_body(
        doc,
        "客户端经 Java WebSocket 接入，服务端桥接讯飞实时 ASR，区分 partial/final；确认文本进入稳定会议记录后，"
        "会后顺序调用转写整理、总结、成员分析、资源推荐等会议类智能体。",
    )
    add_two_figures(doc, [("app-08-meeting-summary.png", "图3-7  AI会议纪要"), ("admin-08-meeting-asr.png", "图3-8  ASR模型配置")])

    add_title(doc, "3.7 题库生成、组卷与考试反馈", 2)
    add_body(
        doc,
        "题库生成支持粘贴资料/上传文件，单次一种题型，生成前结构校验未通过不入库；组卷支持共有/私有/收藏夹选题。"
        "客观题交卷后可按知识点更新掌握度、生成画像候选并重排路径，实现闭环。",
    )
    add_two_figures(doc, [("app-14-question-gen.png", "图3-9  题库生成设置"), ("app-15-paper-select.png", "图3-10  试卷选题")])

    page_break(doc)

    # ==================== Ch4 ====================
    add_title(doc, "第4章  系统实现", 1)
    add_title(doc, "4.1 工程结构与协作契约", 2)
    add_table(
        doc,
        [
            ["模块", "关键目录"],
            ["小程序", "mini_program_app/subpackage_learning|ai|exam|meeting"],
            ["AI服务", "ai-servers/app/multi_agents、learning_workflow"],
            ["Java", "AppBackend/.../controller|service|websocket"],
            ["管理端", "AppWeb/src/pages/ai|questionBank|learning"],
        ],
    )
    add_body(doc, "内部安全：部署注入同一 AI_INTERNAL_TOKEN；Java 请求附带 X-AI-Internal-Token；Python 反向调用转发 Authorization，Java 重新鉴权。")

    add_title(doc, "4.2 小程序实现要点", 2)
    add_body(doc, "学习与 AI 能力以分包接入，不替换既有校园入口。关键页面包括：首页门户、智能助手、画像雷达、PPT/图表生成、题库组卷、会议、课表与课程等。")
    add_two_figures(doc, [("app-01-home.png", "图4-1  首页功能矩阵"), ("app-21-ai-create.png", "图4-2  AI创作与课程入口")])

    add_title(doc, "4.3 Java：画像证据汇总（核心）", 2)
    add_body(doc, "实现类 UserProfileServiceImpl。证据达最低置信度后计算正负加权；相对差小于 0.2 则继续留在候选池；否则融合改分并标记 applied。")
    add_code(
        doc,
        """// UserProfileServiceImpl.java（核心逻辑摘录）
double positiveWeight = validEvidence.stream()
    .filter(e -> e.getSuggestedDelta() > 0)
    .mapToDouble(e -> Math.abs(e.getSuggestedDelta()) * e.getConfidence()).sum();
double negativeWeight = validEvidence.stream()
    .filter(e -> e.getSuggestedDelta() < 0)
    .mapToDouble(e -> Math.abs(e.getSuggestedDelta()) * e.getConfidence()).sum();
if (Math.abs(positiveWeight - negativeWeight) / totalWeight < 0.2) {
    // 正负接近：不改分，继续候选
    return;
}
int appliedDelta = calculateFusionDelta(dimension, rule, averageConfidence, requestedDelta);
// 更新 score/confidence/trend，证据 status -> applied""",
    )
    add_figure(doc, "app-26-profile-radar.png", "图4-3  画像雷达运行结果（基于有效证据更新）", width_cm=7.2)

    add_title(doc, "4.4 Python：Leader 路由（核心）", 2)
    add_code(
        doc,
        """# leader_agent/agent.py（plan 主路径摘录）
forced_plan = self._plan_for_requested_agent(requested_agent, rag_strategy)
if forced_plan: return self._finalize_campus_tool_plan(forced_plan, route_text)
fast_plan = self._plan_high_confidence_service_query(...)
if fast_plan: return self._finalize_campus_tool_plan(fast_plan, route_text)
return self._finalize_campus_tool_plan(self._plan_with_llm(...), route_text)

# Python 课程切入受控工作流
return LeaderPlan(action="run_learning_workflow", route_mode="workflow", ...)""",
    )
    add_two_figures(doc, [("app-27-ai-assistant.png", "图4-4  Leader助手产物卡片"), ("admin-18-agent-test.png", "图4-5  智能体测试调试")])

    add_title(doc, "4.5 Python：课程资源 DAG（核心）", 2)
    add_code(
        doc,
        """# learning_workflow/workflow.py（摘录）
RESOURCE_AGENT_BY_TYPE = {
  "knowledge_note": "textbook_knowledge_agent",
  "mind_map": "diagram_mind_map_agent",
  "practice_set": "python_practice_set_agent",
  "code_lab": "python_code_lab_agent",
  "extended_reading": "extension_reading_agent",
}
MAX_PARALLELISM = 3
# learning_path_agent → 并行资源生成 → resource_review_agent → package""",
    )
    add_two_figures(doc, [("app-16-mindmap.png", "图4-6  思维导图生成配置"), ("app-13-plan-decompose.png", "图4-7  学习计划拆解入口")])

    add_title(doc, "4.6 Java：MaxKB 引用式问答（核心）", 2)
    add_code(
        doc,
        """// KnowledgeChatServiceImpl.java
RetrievalResult retrievalResult = retrieve(retrievalRequest);
chatRequest.setPrompt("回答时优先依据知识库片段；资料不足时明确说明，不要编造。");
chatRequest.setInput(buildAgentInput(question, references));
response.setAnswer(llmService.chat(...).getAnswer());
response.setReferences(references);""",
    )
    add_figure(doc, "admin-21-knowledge.png", "图4-8  知识库与文档治理", width_cm=14.5)

    add_title(doc, "4.7 会议 ASR 与会后链路", 2)
    add_body(doc, "MeetingAsrWebSocketHandler 桥接讯飞结果，区分 isFinal/partial 并向会议端广播 asr_result；会后结果在会议纪要页展示。")
    add_two_figures(doc, [("app-07-meeting-home.png", "图4-9  会议入口"), ("app-09-meeting-detail.png", "图4-10  会议任务与记录")])

    add_title(doc, "4.8 题库、组卷、编程练习与 PPT", 2)
    add_body(doc, "题库生成、试卷选题、编程题库与 PPT 管线分别由对应 Controller/智能体落地；管理端维护公共题库与 Python 题目。")
    add_two_figures(doc, [("app-19-coding-bank.png", "图4-11  编程题库"), ("app-12-ppt-done.png", "图4-12  PPT导出结果")])
    add_figure(doc, "admin-11-question-bank.png", "图4-13  管理端题库列表", width_cm=14.5)

    add_title(doc, "4.9 管理端治理实现", 2)
    add_body(doc, "AppWeb 将 AI 运维与校园业务放在同一控制台：模型配置、智能体绑定、工具监控、观测、知识库、课程与 ASR 配置均可在界面完成。")
    add_figure(doc, "admin-16-agent-bind.png", "图4-14  智能体模型绑定", width_cm=14.5)
    add_figure(doc, "admin-19-tool-monitor.png", "图4-15  工具调用监控", width_cm=14.5)

    page_break(doc)

    # ==================== Ch5 ====================
    add_title(doc, "第5章  测试分析", 1)
    add_title(doc, "5.1 测试方案", 2)
    add_body(
        doc,
        "采用“需求可追踪、核心边界优先、结果可复现、缺口明确披露”原则。"
        "区分：源码存在≠测试通过；离线自动化通过≠真实外部服务通过；静态契约通过≠容器实启/在线压测完成。",
    )

    add_title(doc, "5.2 测试环境", 2)
    add_table(
        doc,
        [
            ["类别", "配置"],
            ["OS", "Windows 10"],
            ["JDK", "21"],
            ["Python", "3.11+（uv）"],
            ["Node", "20+"],
            ["数据与缓存", "MySQL / Redis（Docker）"],
            ["小程序", "微信开发者工具 / 真机预览"],
            ["浏览器", "Chrome"],
        ],
    )

    add_title(doc, "5.3 功能性测试与界面证据", 2)
    add_table(
        doc,
        [
            ["编号", "模块", "验证要点", "界面证据", "结果"],
            ["FT-01", "分端登录", "App/Web角色白名单", "图1-3/图2-系列", "通过"],
            ["FT-02", "Leader助手", "对话、工具产物卡片", "图2-7、图4-4", "通过"],
            ["FT-03", "画像雷达", "七维展示与规则配置", "图2-15、图3-2", "通过"],
            ["FT-04", "图表/PPT", "导图/架构/流程/PPT生成", "图2-10~14", "通过"],
            ["FT-05", "题库组卷", "生成设置与选题", "图2-18、图2-19", "通过"],
            ["FT-06", "编程题库", "难度/标签筛选与进度", "图2-20", "通过"],
            ["FT-07", "会议纪要", "会议入口与AI纪要", "图2-21~23", "通过"],
            ["FT-08", "知识库治理", "管理端知识库配置", "图2-32", "通过"],
            ["FT-09", "智能体治理", "开关/绑定/测试/监控", "图2-35~37、图2-42", "通过"],
            ["FT-10", "校园服务", "活动/论坛/地图/课表等", "图2-1~6", "通过"],
        ],
    )
    add_body(doc, "以上“通过”指当前演示环境下功能可操作且界面可达；真实外部模型质量、MaxKB 召回率、ASR 准确率与压测指标需在正式环境另行留证，本节不虚构数值。")

    add_title(doc, "5.4 自动化与边界", 2)
    add_body(doc, "仓库具备 Java/Python/AppWeb 离线测试与质量门禁脚本。画像冲突门禁、工作流 partial/重试、考试反馈幂等等边界有自动化覆盖；提交前建议在最终 SHA 重跑门禁并保留终端截图。")

    page_break(doc)

    # ==================== Ch6 ====================
    add_title(doc, "第6章  作品总结", 1)
    add_title(doc, "6.1 作品特色与创新点", 2)
    add_bullet(doc, "证据驱动的七维画像慢更新：候选池 + 置信度加权 + 正负冲突门禁。")
    add_bullet(doc, "双拓扑多智能体：校园 Leader 单目标路由 + 课程 typed DAG；Catalog 可扩展。")
    add_bullet(doc, "检索与生成职责分离：MaxKB hit-test + 本系统 Agent 回答与 citations。")
    add_bullet(doc, "多模态学习资源：讲解/导图/练习/代码/拓展 + 独立 PPT 管线，支持复审与导出。")
    add_bullet(doc, "学—练—测闭环：题库生成、组卷、编程练习与画像/路径反馈衔接。")
    add_bullet(doc, "智能会议：实时 ASR 桥接 + 会后纪要智能体链。")
    add_bullet(doc, "校园能力增量接入：AI 学习与校园服务同端共存，管理端一体化治理。")

    add_title(doc, "6.2 作品展望", 2)
    add_bullet(doc, "完成真实课程知识库合法导出、空环境恢复与引用映射验收。")
    add_bullet(doc, "在真实 endpoint 完成事实评测与负载测试并保留原始报告。")
    add_bullet(doc, "强化 Web 统一 RBAC、会议 WebSocket Origin/限流，推进密码自适应哈希。")
    add_bullet(doc, "向更多课程复用 DAG 与画像协议，扩展辅导形态，同时保持业务事实由 Java 治理。")

    add_title(doc, "6.3 结束语", 2)
    add_body(
        doc,
        "AgentA3 的价值在于把大模型能力放进可审计的业务状态机：谁能调用、证据如何生效、资源如何交付、"
        "考试如何反馈，都有明确工程边界。后续重点从功能演示转向真实知识、真实评测与生产加固，"
        "使作品从可演示系统进一步成为可核验、可运维的校园学习智能体平台。",
    )

    # Appendix: full figure list
    page_break(doc)
    add_title(doc, "附录  截图清单", 1)
    add_body(doc, "本报告正文已嵌入全部 49 张功能截图。清单如下，便于对照源文件。")
    rows = [["序号", "文件名", "说明"]]
    catalog = [
        ("admin-01-users.png", "用户与角色管理"),
        ("admin-02-activities.png", "活动管理"),
        ("admin-03-forum.png", "论坛帖子管理"),
        ("admin-04-canteen.jpg", "食堂管理"),
        ("admin-05-map.jpg", "校园标点管理"),
        ("admin-06-secondhand.png", "闲置物品管理"),
        ("admin-07-merchants.png", "商家管理"),
        ("admin-08-meeting-asr.png", "会议语音模型配置"),
        ("admin-09-courses.png", "校园课程管理"),
        ("admin-10-python-problems.png", "Python题库管理"),
        ("admin-11-question-bank.png", "题库题目管理"),
        ("admin-12-ai-chat.png", "管理端智能助手"),
        ("admin-13-model-config.png", "AI模型配置"),
        ("admin-14-agent-tools.png", "智能体工具开关"),
        ("admin-15-agent-formats.png", "智能体文件格式设置"),
        ("admin-16-agent-bind.png", "智能体模型绑定"),
        ("admin-17-question-mapping.png", "智能体题型映射"),
        ("admin-18-agent-test.png", "智能体测试调试"),
        ("admin-19-tool-monitor.png", "工具调用监控"),
        ("admin-20-observability.png", "Langfuse观测配置"),
        ("admin-21-knowledge.png", "AI知识库管理"),
        ("admin-22-profile-rules.png", "画像维度规则配置"),
        ("app-01-home.png", "小程序首页"),
        ("app-02-activities.png", "校园活动"),
        ("app-03-secondhand.png", "闲置交易"),
        ("app-04-forum.png", "校园论坛"),
        ("app-05-coupons.png", "优惠券"),
        ("app-06-map.png", "校园地图"),
        ("app-07-meeting-home.png", "会议首页"),
        ("app-08-meeting-summary.png", "AI会议纪要"),
        ("app-09-meeting-detail.png", "会议详情"),
        ("app-10-smart-writing.png", "智能写作"),
        ("app-11-ppt-template.png", "PPT模板选择"),
        ("app-12-ppt-done.png", "PPT生成完成"),
        ("app-13-plan-decompose.png", "学习计划拆解"),
        ("app-14-question-gen.png", "题库生成"),
        ("app-15-paper-select.png", "试卷选题"),
        ("app-16-mindmap.png", "AI思维导图"),
        ("app-17-architecture.png", "AI架构图"),
        ("app-18-flowchart.png", "AI流程图"),
        ("app-19-coding-bank.png", "编程题库"),
        ("app-20-pdf2word.png", "PDF转Word"),
        ("app-21-ai-create.png", "AI创作门户"),
        ("app-22-watermark.png", "水印工具"),
        ("app-23-schedule.png", "我的课表"),
        ("app-24-my-courses.png", "我的课程"),
        ("app-25-ai-history.png", "AI会话历史"),
        ("app-26-profile-radar.png", "个人画像雷达"),
        ("app-27-ai-assistant.png", "智能助手对话"),
    ]
    for i, (fn, desc) in enumerate(catalog, 1):
        rows.append([str(i), fn, desc])
    add_table(doc, rows)

    doc.save(str(OUT))
    try:
        import shutil

        shutil.copy2(OUT, OUT_DESKTOP)
    except Exception as exc:
        print("desktop copy failed:", exc)
    print("saved", OUT, "size", OUT.stat().st_size)


if __name__ == "__main__":
    build()
